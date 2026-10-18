import io
import csv

from fastapi.testclient import TestClient
from mission_planner.main import app

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


client = TestClient(app)


def _create_and_prepare_project():
    # create project, generate plan and checklists
    r = client.post("/api/projects", json={"idea": "automated test project", "domain": "space_mission"})
    assert r.status_code == 200
    pid = r.json()["project_id"]
    # generate a deterministic plan
    r = client.post(f"/api/projects/{pid}/generate-plan", json={})
    assert r.status_code == 200
    # generate checklists (deterministic fallback used in test env)
    r = client.post(f"/api/projects/{pid}/generate-checklists")
    assert r.status_code == 200
    return pid


def test_stream_csv_and_filters():
    pid = _create_and_prepare_project()

    # full CSV
    r = client.get(f"/api/projects/{pid}/export_stream.csv")
    assert r.status_code == 200
    assert "text/csv" in r.headers.get("content-type", "")
    text = r.content.decode("utf-8")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    assert rows, "CSV should contain header and rows"
    header = rows[0]
    assert "id" in header and "phase" in header

    # filtered CSV: Phase B
    r2 = client.get(f"/api/projects/{pid}/export_stream.csv", params={"phase": "Phase B"})
    assert r2.status_code == 200
    text2 = r2.content.decode("utf-8")
    reader2 = csv.reader(io.StringIO(text2))
    rows2 = list(reader2)
    # every row after header must have Phase B in the phase column (index 1)
    for row in rows2[1:]:
        assert row[1] == "Phase B"


def test_stream_xlsx_sheets():
    if load_workbook is None:
        # openpyxl not available in test environment
        return

    pid = _create_and_prepare_project()
    r = client.get(f"/api/projects/{pid}/export_stream.xlsx", params={"sheets": "both"})
    assert r.status_code == 200
    ct = r.headers.get("content-type", "")
    assert "spreadsheetml" in ct or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in ct

    wb = load_workbook(filename=io.BytesIO(r.content))
    # check main sheet exists and there is at least one phase sheet
    assert "All Checklists" in wb.sheetnames
    assert any(name.startswith("Phase") or name.startswith("Pre-Phase") for name in wb.sheetnames)
